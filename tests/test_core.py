"""Core regression tests. Run: python -m pytest tests -q

NOTE: the decryptor test builds a synthetic crypt14-style file using this
project's own assumed framing — it proves the zlib+GCM pipeline and tag
handling are self-consistent, NOT that the framing matches a real
WhatsApp release. Validate against a real backup before casework."""
import os, sqlite3, sys, zlib
from datetime import timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from Crypto.Cipher import AES
from openpyxl import load_workbook

from modules import decryptor, custody
from modules.db_parser import WaDatabase, Chat, Message, CallRecord
from modules import export


def _varint(n):
    out = bytearray()
    while True:
        b = n & 0x7F; n >>= 7
        out.append(b | (0x80 if n else 0))
        if not n: return bytes(out)


def make_sqlite(path, n_messages=1200):
    c = sqlite3.connect(path)
    c.executescript("""
    CREATE TABLE jid(_id INTEGER PRIMARY KEY, raw_string TEXT);
    CREATE TABLE chat(_id INTEGER PRIMARY KEY, jid_row_id INT);
    CREATE TABLE message(_id INTEGER PRIMARY KEY, chat_row_id INT, from_me INT,
                         text_data TEXT, timestamp INT, message_type INT);
    CREATE TABLE call_log(_id INTEGER PRIMARY KEY, jid_row_id INT, from_me INT,
                          video_call INT, duration INT, timestamp INT, call_result INT);
    CREATE TABLE message_thumbnails(message_row_id INT, thumbnail BLOB);
    INSERT INTO jid VALUES (1,'1111@s.whatsapp.net');
    INSERT INTO chat VALUES (1,1);
    """)
    for i in range(n_messages):
        c.execute("INSERT INTO message VALUES (?,?,?,?,?,?)",
                  (i + 1, 1, i % 2, f"msg {i}", 1_700_000_000_000 + i * 1000, 15 if i == 3 else 0))
    c.execute("INSERT INTO message VALUES (?,?,?,?,?,?)",
              (n_messages + 1, 1, 0, "=HYPERLINK(\"http://x\",\"y\")", 1_700_000_000_000, 0))
    c.execute("INSERT INTO message_thumbnails VALUES (1, x'FFD8FF')")
    c.execute("INSERT INTO call_log VALUES (1,1,0,0,0,1700000000000,5)")   # incoming, no talk
    c.execute("INSERT INTO call_log VALUES (2,1,1,0,0,1700000001000,2)")   # outgoing, no talk
    c.execute("INSERT INTO call_log VALUES (3,1,0,1,42,1700000002000,1)")  # incoming, answered
    c.commit(); c.close()


def test_no_message_limit_and_utc_and_flags(tmp_path):
    db_path = str(tmp_path / "m.db"); make_sqlite(db_path)
    db = WaDatabase(db_path)
    chats = db.get_chats()
    msgs = chats[0].messages
    assert len(msgs) == 1201                      # was silently capped at 500
    assert msgs[0].timestamp.tzinfo == timezone.utc
    by_text = {m.text: m for m in msgs}
    assert by_text["msg 3"].is_deleted and not by_text["msg 4"].is_deleted
    assert by_text["msg 0"].thumbnail_blob == b"\xff\xd8\xff"
    assert len(db.get_chats(limit=10)[0].messages) == 10
    types = [c.call_type for c in sorted(db.get_calls(), key=lambda c: c.timestamp)]
    assert types == ["MISSED", "UNANSWERED", "IN"]
    assert db.get_calls()[0].result_code is not None


@pytest.mark.parametrize("footer", [0, 16, 20])
def test_decrypt_roundtrip(tmp_path, footer):
    key = os.urandom(32); iv = os.urandom(16)
    db_path = str(tmp_path / "plain.db"); make_sqlite(db_path, 50)
    plain = open(db_path, "rb").read()
    cipher = AES.new(key, AES.MODE_GCM, nonce=iv)
    ct = cipher.encrypt(zlib.compress(plain)); tag = cipher.digest()
    proto = b"\x0a" + _varint(16) + iv
    blob = b"\x00" + _varint(len(proto)) + proto + ct + tag + os.urandom(footer)
    enc = tmp_path / "msgstore.db.crypt14"; enc.write_bytes(blob)
    keyfile = tmp_path / "key"; keyfile.write_bytes(os.urandom(126) + key)
    res = decryptor.decrypt_backup(str(keyfile), str(enc), str(tmp_path / "out"))
    assert open(res.output_path, "rb").read() == plain
    WaDatabase(res.output_path).get_chats()       # actually opens as SQLite
    ok, _ = custody.verify_log(str(tmp_path / "out"))
    assert ok


def test_decrypt_wrong_key_leaves_no_output(tmp_path):
    key = os.urandom(32); iv = os.urandom(16)
    cipher = AES.new(key, AES.MODE_GCM, nonce=iv)
    ct = cipher.encrypt(zlib.compress(b"hello" * 100)); tag = cipher.digest()
    proto = b"\x0a" + _varint(16) + iv
    enc = tmp_path / "x.crypt14"; enc.write_bytes(b"\x00" + _varint(len(proto)) + proto + ct + tag)
    keyfile = tmp_path / "key"; keyfile.write_bytes(os.urandom(32))
    with pytest.raises(decryptor.DecryptError):
        decryptor.decrypt_backup(str(keyfile), str(enc), str(tmp_path / "out"))
    assert not os.path.exists(tmp_path / "out" / "msgstore.decrypted.db")
    assert not os.path.exists(tmp_path / "out" / "msgstore.decrypted.db.partial")


def test_crypt15_hex_key_derivation(tmp_path):
    kf = tmp_path / "k.txt"; kf.write_text("ab" * 32)
    assert decryptor.load_key(str(kf), crypt15=True) != bytes.fromhex("ab" * 32)
    assert decryptor.load_key(str(kf), crypt15=False) == bytes.fromhex("ab" * 32)


def test_xlsx_formula_injection_and_name_collision(tmp_path):
    db_path = str(tmp_path / "m.db"); make_sqlite(db_path, 5)
    chats = WaDatabase(db_path).get_chats()
    twin = Chat(jid="2222@s.whatsapp.net", display_name=chats[0].display_name, messages=[
        Message(chat_jid="2222@s.whatsapp.net", from_me=True, text="hi", timestamp=None)])
    out = str(tmp_path / "exp")
    written = export.export_all(chats + [twin], [], out, ["txt", "xlsx", "html", "reconstruct"])
    assert len(set(written["txt"])) == 2          # no overwrite
    ws = load_workbook(written["xlsx"][0]).active
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.data_type != "f" for c in cells)
    assert "UTC" in open(written["txt"][0], encoding="utf-8").read()
    ok, msg = custody.verify_log(out); assert ok, msg
    assert os.path.exists(os.path.join(out, "case_report.txt"))


def test_custody_tamper_detected(tmp_path):
    f = tmp_path / "a.bin"; f.write_bytes(b"x")
    for _ in range(3): custody.record(str(tmp_path), "t", str(f))
    assert custody.verify_log(str(tmp_path))[0]
    p = tmp_path / custody.LOG_NAME
    lines = p.read_text().splitlines(); del lines[1]
    p.write_text("\n".join(lines) + "\n")
    assert not custody.verify_log(str(tmp_path))[0]


def test_record_tree_manifest(tmp_path):
    media = tmp_path / "Media"; (media / "a").mkdir(parents=True)
    (media / "a" / "x.jpg").write_bytes(b"1"); (media / "y.mp4").write_bytes(b"22")
    rec = custody.record_tree(str(tmp_path), "pull-media", str(media))
    assert "2 files" in rec["note"] and rec["sha256"]
    lines = (tmp_path / "pull-media_manifest.sha256").read_text().splitlines()
    assert len(lines) == 2 and any(l.endswith("  a/x.jpg") for l in lines)
