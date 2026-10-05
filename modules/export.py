"""
Export
------
Renders parsed chats/calls to court-referenceable formats:
  - TXT   (plain chat-format transcript)
  - XLSX  (tabular, sortable — good for indexing/searching)
  - HTML  (WhatsApp-style bubble rendering, for readability in a report)
  - PDF   (optional; requires the pdf skill / a PDF engine if wkhtmltopdf
           or a headless browser isn't available on the examiner's machine
           — this module produces the HTML and leaves final PDF rendering
           to whatever engine you wire up, since that's environment-specific)
"""

import os
import html
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from .db_parser import Chat, CallRecord, MIN_TS, fmt_ts
from . import custody

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
VIDEO_EXTS = {".mp4", ".3gp", ".mov", ".mkv"}
AUDIO_EXTS = {".opus", ".mp3", ".m4a", ".aac", ".wav"}


def _append_safe(ws, row):
    """Append a row while neutralising spreadsheet formula injection.
    openpyxl turns any string starting with '=' into a FORMULA cell, and
    message text is attacker-controlled. After appending, force those
    cells back to plain-text type so the content is preserved verbatim
    and never evaluated."""
    ws.append(list(row))
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _unique_name(base: str, ext: str, used: set) -> str:
    """Return a file name that has not been used in this export run, so two
    chats with the same display name don't overwrite each other."""
    name = f"{base}{ext}"
    n = 2
    while name.lower() in used:
        name = f"{base}_{n}{ext}"
        n += 1
    used.add(name.lower())
    return name


def export_chat_txt(chat: Chat, output_path: str):
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(f"Chat with: {chat.display_name} ({chat.jid})\n")
        f.write(f"Exported: {datetime.now(timezone.utc).isoformat(timespec='seconds')} (all message times are UTC)\n")
        f.write("=" * 60 + "\n\n")
        for m in chat.messages:
            who = "Me" if m.from_me else chat.display_name
            ts = fmt_ts(m.timestamp, seconds=True)
            f.write(f"[{ts}] {who}: {m.text}\n")


def export_chats_xlsx(chats: list[Chat], output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Messages"
    ws.append(["Chat", "From", "Message", "Timestamp (UTC)", "Media Type", "Deleted for everyone"])
    for chat in chats:
        for m in chat.messages:
            _append_safe(ws, [
                chat.display_name,
                "Me" if m.from_me else chat.display_name,
                m.text,
                m.timestamp.isoformat() if m.timestamp else "",
                m.media_type or "",
                "Yes" if m.is_deleted else "",
            ])
    wb.save(output_path)


def export_calls_xlsx(calls: list[CallRecord], output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Call Log"
    ws.append(["Contact/Number", "Type", "Video", "Duration (s)", "Timestamp (UTC)", "Raw result code"])
    for c in calls:
        _append_safe(ws, [
            c.phone_number,
            c.call_type,
            "Yes" if c.is_video else "No",
            c.duration_seconds,
            c.timestamp.isoformat() if c.timestamp else "",
            c.result_code if c.result_code is not None else "",
        ])
    wb.save(output_path)


_HTML_TEMPLATE = """<!DOCTYPE html>
<html><head><meta charset="utf-8">
<title>Chat Export — {chat_name}</title>
<style>
  body {{ background:#e5ddd5; font-family: -apple-system, Segoe UI, sans-serif; margin:0; padding:20px; }}
  .header {{ background:#075e54; color:white; padding:14px 20px; border-radius:8px 8px 0 0; }}
  .chat-container {{ max-width:700px; margin:0 auto; background:white; border-radius:0 0 8px 8px; padding:16px; }}
  .bubble {{ max-width:70%; padding:8px 12px; border-radius:8px; margin:6px 0; font-size:14px; line-height:1.4; clear:both; }}
  .me {{ background:#dcf8c6; float:right; }}
  .them {{ background:#ffffff; border:1px solid #eee; float:left; }}
  .meta {{ font-size:11px; color:#888; margin-top:4px; }}
  .clearfix {{ clear:both; }}
  .media-missing {{ font-size:12px; color:#a55; background:#fff3f3; border:1px dashed #e0b4b4;
                     border-radius:6px; padding:6px 8px; margin-top:6px; }}
  .call-marker {{ clear:both; text-align:center; margin:14px auto; }}
  .call-marker span {{ display:inline-block; background:#e3f2fd; color:#2a6f97; font-size:12px;
                        padding:5px 14px; border-radius:14px; }}
  img, video {{ max-width:280px; border-radius:8px; display:block; margin-top:6px; }}
  audio {{ display:block; margin-top:6px; max-width:280px; }}
</style></head>
<body>
<div class="header"><strong>{chat_name}</strong><br><span style="font-size:12px;">{jid}</span></div>
<div class="chat-container">
{messages}
<div class="clearfix"></div>
</div>
</body></html>
"""

_BUBBLE_TEMPLATE = """<div class="bubble {cls}">{text}<div class="meta">{ts}</div></div><div class="clearfix"></div>"""
_CALL_MARKER_TEMPLATE = """<div class="call-marker"><span>{icon} {label}{duration}</span><div class="meta">{ts}</div></div>"""


def _media_snippet(local_path, relative_path, base_dir=None) -> str:
    """Embed the actual file if we resolved a local copy of it; otherwise
    say plainly that the reference exists but wasn't recovered locally."""
    if not relative_path:
        return ""
    if not local_path:
        return (f'<div class="media-missing">📎 media referenced '
                 f'({html.escape(os.path.basename(relative_path))}) — not recovered locally</div>')
    # Prefer a path relative to the report so the export folder can be
    # moved/zipped together with its media; fall back to an absolute
    # file:// URI (e.g. media on a different drive).
    resolved = Path(local_path).resolve()
    uri = resolved.as_uri()
    if base_dir:
        try:
            uri = Path(os.path.relpath(resolved, base_dir)).as_posix().replace(" ", "%20")
        except ValueError:
            pass
    ext = os.path.splitext(local_path)[1].lower()
    if ext in IMAGE_EXTS:
        return f'<img src="{uri}" alt="media">'
    elif ext in VIDEO_EXTS:
        return f'<video controls><source src="{uri}"></video>'
    elif ext in AUDIO_EXTS:
        return f'<audio controls><source src="{uri}"></audio>'
    else:
        return f'<a href="{uri}">📎 {html.escape(os.path.basename(local_path))}</a>'


def _call_marker(call: CallRecord) -> str:
    icon = "📹" if call.is_video else "📞"
    label = {"IN": "Incoming call", "OUT": "Outgoing call", "MISSED": "Missed call"}.get(
        call.call_type, call.call_type
    )
    duration = ""
    if call.duration_seconds:
        duration = f" · {call.duration_seconds // 60}:{call.duration_seconds % 60:02d}"
    ts = fmt_ts(call.timestamp, empty="")
    return _CALL_MARKER_TEMPLATE.format(icon=icon, label=label, duration=duration, ts=ts)


def export_chat_html(chat: Chat, output_path: str, media_index=None):
    rows = []
    for m in chat.messages:
        cls = "me" if m.from_me else "them"
        ts = fmt_ts(m.timestamp, empty="")
        text_html = html.escape(m.text).replace("\n", "<br>") if m.text else ""
        local_path = media_index.resolve(m.media_relative_path) if media_index else None
        media_html = _media_snippet(local_path, m.media_relative_path, os.path.dirname(os.path.abspath(output_path)))
        rows.append(_BUBBLE_TEMPLATE.format(cls=cls, text=text_html + media_html, ts=ts))
    page = _HTML_TEMPLATE.format(
        chat_name=html.escape(chat.display_name),
        jid=html.escape(chat.jid),
        messages="\n".join(rows),
    )
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(page)


def export_reconstructed_chat(chat: Chat, calls: list[CallRecord], output_path: str, media_index=None):
    """
    'Recreate the whole chat': merges every message AND every call with
    this contact into one chronological timeline, with media embedded
    inline wherever a local copy was resolved (extracted Media folder or
    a file recovered via the filesystem scanner). This is the full
    activity reconstruction, not just the text transcript.
    """
    calls_for_chat = [c for c in calls if c.contact_jid == chat.jid]

    events = []
    for m in chat.messages:
        events.append((m.timestamp or MIN_TS, "message", m))
    for c in calls_for_chat:
        events.append((c.timestamp or MIN_TS, "call", c))
    events.sort(key=lambda e: e[0])

    rows = []
    for ts, kind, payload in events:
        if kind == "message":
            m = payload
            cls = "me" if m.from_me else "them"
            ts_str = fmt_ts(m.timestamp, empty="")
            text_html = html.escape(m.text).replace("\n", "<br>") if m.text else ""
            local_path = media_index.resolve(m.media_relative_path) if media_index else None
            media_html = _media_snippet(local_path, m.media_relative_path, os.path.dirname(os.path.abspath(output_path)))
            rows.append(_BUBBLE_TEMPLATE.format(cls=cls, text=text_html + media_html, ts=ts_str))
        else:
            rows.append(_call_marker(payload))

    page = _HTML_TEMPLATE.format(
        chat_name=html.escape(chat.display_name) + " — full reconstruction",
        jid=html.escape(chat.jid),
        messages="\n".join(rows),
    )
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(page)


def export_extras_xlsx(db, output_path: str) -> int:
    """One workbook with reactions, edits, group participants, starred
    messages and status updates read from the open WaDatabase. Returns the
    total number of rows written."""
    wb = Workbook()
    wb.remove(wb.active)
    total = 0

    def sheet(name, header, rows):
        nonlocal total
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            _append_safe(ws, r)
        total += len(rows)

    sheet("Reactions", ["Chat", "Reactor", "Emoji", "On message", "Time (UTC)"],
          [[r.chat_jid, r.reactor_jid, r.emoji, r.target_text, r.timestamp.isoformat() if r.timestamp else ""]
           for r in db.get_reactions()])
    sheet("Edits", ["Chat", "Current text", "Edited (UTC)", "Originally sent (UTC)"],
          [[e.chat_jid, e.current_text, e.edited_at.isoformat() if e.edited_at else "",
            e.original_sent_at.isoformat() if e.original_sent_at else ""] for e in db.get_edits()])
    sheet("Group participants", ["Group", "Group name", "Member", "Role"],
          [[p.group_jid, p.group_name, p.member_jid, p.role] for p in db.get_group_participants()])
    sheet("Starred", ["Chat", "Text", "Time (UTC)"],
          [[m.chat_jid, m.text, m.timestamp.isoformat() if m.timestamp else ""] for m in db.get_starred()])
    sheet("Status updates", ["From me", "Text", "Media", "Time (UTC)"],
          [["Yes" if m.from_me else "No", m.text, m.media_relative_path or "",
            m.timestamp.isoformat() if m.timestamp else ""] for m in db.get_status_updates()])
    wb.save(output_path)
    return total


def export_all(chats: list[Chat], calls: list[CallRecord], output_dir: str,
               formats: list[str], media_index=None):
    """
    formats: any combination of "txt", "xlsx", "html", "calls_xlsx", "reconstruct"
    media_index: an optional media_resolver.MediaIndex used to embed real
                 media inline for "html" and "reconstruct" formats.
    Returns a dict of format -> list of written file paths.
    """
    os.makedirs(output_dir, exist_ok=True)
    written = {fmt: [] for fmt in formats}

    used_names: set = set()
    for chat in chats:
        base = "".join(c for c in chat.display_name if c.isalnum() or c in " _-")[:50] or "chat"
        # one stem per chat, shared by all its formats, unique across chats
        stem = _unique_name(base, "", used_names)
        if "txt" in formats:
            p = os.path.join(output_dir, f"{stem}.txt")
            export_chat_txt(chat, p)
            written["txt"].append(p)
        if "html" in formats:
            p = os.path.join(output_dir, f"{stem}.html")
            export_chat_html(chat, p, media_index=media_index)
            written["html"].append(p)
        if "reconstruct" in formats:
            p = os.path.join(output_dir, f"{stem}.reconstructed.html")
            export_reconstructed_chat(chat, calls, p, media_index=media_index)
            written["reconstruct"].append(p)

    if "xlsx" in formats:
        p = os.path.join(output_dir, "all_messages.xlsx")
        export_chats_xlsx(chats, p)
        written["xlsx"].append(p)

    if "calls_xlsx" in formats and calls:
        p = os.path.join(output_dir, "call_log.xlsx")
        export_calls_xlsx(calls, p)
        written["calls_xlsx"].append(p)

    # Chain of custody: hash and log every file written.
    for fmt, paths in written.items():
        for p in paths:
            custody.record(output_dir, f"export-{fmt}", p)
    custody.write_case_report(output_dir)
    return written
